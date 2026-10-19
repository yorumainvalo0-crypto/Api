"""Runs real executions in the sandbox image (needs Docker and portal-sandbox:dev)."""

from __future__ import annotations

import io
import os
import re
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from runner.app import create_app
from runner.sandbox import Limits, Sandbox
from runner.storage import Store

IMAGE = os.environ.get("SANDBOX_IMAGE", "portal-sandbox:dev")
ID = re.compile(r"^[A-Za-z0-9_-]{21}$")


@pytest.fixture(scope="module")
def client(tmp_path_factory: pytest.TempPathFactory) -> TestClient:
    root = tmp_path_factory.mktemp("runner")
    store = Store(root / "store")
    sandbox = Sandbox(store, root / "work", IMAGE, Limits(timeout_seconds=5))
    return TestClient(create_app(store, sandbox))


def run(client: TestClient, lang: str, code: str, files: list[dict] | None = None) -> dict:
    res = client.post("/v1/exec", json={"lang": lang, "code": code, "files": files or [], "user_id": "u1", "codeSessionKey": "x"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert ID.match(body["session_id"])
    return body


def download(client: TestClient, ref: dict) -> bytes:
    res = client.get(f"/v1/download/{ref['storage_session_id']}/{ref['id']}", params={"kind": "user", "id": "u1"})
    assert res.status_code == 200
    return res.content


def test_generates_office_documents_and_pdf(client: TestClient) -> None:
    code = """
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

d = Document(); d.add_heading('Angebot', 0); d.save('angebot.docx')
wb = Workbook(); wb.active['A1'] = 'Summe'; wb.active['B1'] = 42; wb.save('zahlen.xlsx')
p = Presentation(); p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = 'Hallo'; p.save('folien.pptx')
c = canvas.Canvas('bericht/report.pdf') if False else None
import os; os.makedirs('bericht', exist_ok=True)
c = canvas.Canvas('bericht/report.pdf'); c.drawString(72, 720, 'Bericht'); c.save()
print('fertig')
"""
    body = run(client, "py", code)
    assert body["stdout"].strip() == "fertig"
    names = sorted(f["name"] for f in body["files"])
    assert names == ["angebot.docx", "bericht/report.pdf", "folien.pptx", "zahlen.xlsx"]
    refs = {f["name"]: f for f in body["files"]}
    assert all(ID.match(f["id"]) and f["storage_session_id"] == body["session_id"] for f in body["files"])
    for name in ("angebot.docx", "zahlen.xlsx", "folien.pptx"):
        assert zipfile.is_zipfile(io.BytesIO(download(client, refs[name])))
    assert download(client, refs["bericht/report.pdf"]).startswith(b"%PDF")


def test_uploaded_files_are_mounted_and_unchanged_inputs_not_reported(client: TestClient) -> None:
    up = client.post("/v1/upload", files={"file": ("daten.csv", b"a,b\n1,2\n3,4\n")}, data={"kind": "user", "id": "u1"})
    assert up.status_code == 200
    payload = up.json()
    assert payload["message"] == "success"
    ref = {"id": payload["files"][0]["fileId"], "storage_session_id": payload["storage_session_id"], "name": "daten.csv"}

    body = run(client, "bash", "wc -l < daten.csv; cp daten.csv kopie.csv")
    assert "No such file" in body["stderr"]  # no inputs passed in this call

    body = run(client, "bash", "wc -l < daten.csv; echo '5,6' >> daten.csv; cat daten.csv > summe.txt", [ref])
    assert body["stdout"].strip() == "3"
    names = sorted(f["name"] for f in body["files"])
    # daten.csv changed (new id), summe.txt is new
    assert names == ["daten.csv", "summe.txt"]
    changed = next(f for f in body["files"] if f["name"] == "daten.csv")
    assert changed["id"] != ref["id"]

    body = run(client, "bash", "cat daten.csv > /dev/null", [ref])
    assert body["files"] == []


def test_deleted_inputs_are_reported(client: TestClient) -> None:
    up = client.post("/v1/upload", files={"file": ("weg.txt", b"x")}).json()
    ref = {"id": up["files"][0]["fileId"], "storage_session_id": up["storage_session_id"], "name": "weg.txt"}
    body = run(client, "bash", "rm weg.txt", [ref])
    assert body["deleted_files"] == ["weg.txt"]


def test_no_network(client: TestClient) -> None:
    body = run(client, "py", "import urllib.request\nurllib.request.urlopen('http://1.1.1.1', timeout=3)")
    assert "Error" in body["stderr"] or "error" in body["stderr"]
    assert "Exit-Code" in body["stderr"]


def test_timeout_kills_the_run(client: TestClient) -> None:
    body = run(client, "bash", "echo start; sleep 30; echo never")
    assert body["status"] == "TO"
    assert "Zeitlimit" in body["stderr"]
    assert "never" not in body["stdout"]


def test_cannot_write_outside_mnt_data_or_escalate(client: TestClient) -> None:
    body = run(client, "bash", "touch /etc/hacked; id -u; touch /tmp/scratch && echo tmp-ok")
    assert "Read-only file system" in body["stderr"] or "Permission denied" in body["stderr"]
    assert body["stdout"].split() == ["1000", "tmp-ok"]
    assert body["files"] == []


def test_path_traversal_in_input_names_is_ignored(client: TestClient) -> None:
    up = client.post("/v1/upload", files={"file": ("ok.txt", b"x")}).json()
    ref = {"id": up["files"][0]["fileId"], "storage_session_id": up["storage_session_id"], "name": "../../escape.txt"}
    body = run(client, "bash", "ls -A", [ref])
    assert body["stdout"].strip() == ""


def test_read_only_skill_files(client: TestClient) -> None:
    batch = client.post(
        "/v1/upload/batch",
        files=[("file", ("skills/pdf/SKILL.md", b"# PDF")), ("file", ("skills/pdf/scripts/make.py", b"print('skill')"))],
        data={"kind": "skill", "id": "s1", "version": "1", "read_only": "true"},
    ).json()
    assert batch["message"] == "success" and batch["succeeded"] == 2
    assert [f["filename"] for f in batch["files"]] == ["skills/pdf/SKILL.md", "skills/pdf/scripts/make.py"]
    refs = [
        {"id": f["fileId"], "storage_session_id": batch["storage_session_id"], "name": f["filename"], "kind": "skill"}
        for f in batch["files"]
    ]
    body = run(client, "bash", "python3 skills/pdf/scripts/make.py; echo x >> skills/pdf/scripts/make.py", refs)
    assert body["stdout"].strip() == "skill"
    assert "Permission denied" in body["stderr"]
    assert body["files"] == []


def test_metadata_listing_and_delete(client: TestClient) -> None:
    up = client.post("/v1/upload", files={"file": ("info.txt", b"hallo")}).json()
    sid, fid = up["storage_session_id"], up["files"][0]["fileId"]
    info = client.get(f"/v1/sessions/{sid}/objects/{fid}").json()
    assert info["originalFilename"] == "info.txt" and info["size"] == 5 and "GMT" in info["lastModified"]
    listing = client.get(f"/v1/files/{sid}", params={"detail": "summary"}).json()
    assert listing[0]["fileId"] == fid
    assert client.delete(f"/v1/files/{sid}/{fid}").status_code == 200
    assert client.get(f"/v1/download/{sid}/{fid}").status_code == 404
    assert client.get("/v1/download/bad/id").status_code == 400


def test_unsupported_language(client: TestClient) -> None:
    body = run(client, "rs", "fn main() {}")
    assert "nicht unterstützt" in body["stderr"]


def test_pdf_with_umlauts_using_the_skill_font(client: TestClient) -> None:
    code = """
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
pdfmetrics.registerFont(TTFont('DejaVu', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
c = canvas.Canvas('/mnt/data/rechnung.pdf', pagesize=A4)
c.setFont('DejaVu', 12); c.drawString(72, 760, 'Größe: 1.234,50 €'); c.save()
from pypdf import PdfReader
print(PdfReader('/mnt/data/rechnung.pdf').pages[0].extract_text())
"""
    body = run(client, "py", code)
    assert "Größe: 1.234,50 €" in body["stdout"], body["stderr"]
    assert [f["name"] for f in body["files"]] == ["rechnung.pdf"]
